"""Extraction of the distributor price list (reference data, rules R401-R405).

Each sheet holds blocks: a heading in column A names a product family, a row with
``PART NUMBER | DESCRIPCIÓN | PRECIO`` opens its table, and item rows follow. The list holds no
personal data.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from pcsforecast.text import normalize_text

HEADING_ROW = ("PART NUMBER", "DESCRIPCIÓN", "PRECIO")
PART_COLUMN, DESCRIPTION_COLUMN, PRICE_COLUMN = 1, 2, 3
COLUMNS = [
    "item_id",
    "source_sha256",
    "sheet",
    "row",
    "family",
    "part_number",
    "description",
    "price",
    "has_price",
    "part_number_repeated",
    "extra_cells",
]


def clean_part_number(value: object) -> str | None:
    """Part number without surrounding parentheses or spaces (rule R402)."""
    text = normalize_text(value)
    if text is None:
        return None
    return text.strip("() ") or None


def _is_heading_row(cells: Sequence[object]) -> bool:
    found = tuple(normalize_text(c) for c in cells[PART_COLUMN : PRICE_COLUMN + 1])
    return found == HEADING_ROW


def _family_heading(cells: Sequence[object]) -> str | None:
    """Text alone in column A (rule R401)."""
    filled = [c for c in cells if c not in (None, "")]
    if len(filled) == 1 and cells and isinstance(cells[0], str):
        return normalize_text(cells[0])
    return None


def _price(value: object) -> float | None:
    if isinstance(value, int | float) and not isinstance(value, bool):
        return float(value)
    return None


def _extra_cells(cells: Sequence[object]) -> str | None:
    """Values right of the price column, as ``<column letter>=<value>`` (rule R404)."""
    extras = [
        f"{openpyxl.utils.get_column_letter(index + 1)}={value}"
        for index, value in enumerate(cells)
        if index > PRICE_COLUMN and value not in (None, "")
    ]
    return ";".join(extras) or None


def sheet_items(sheet: str, rows: Iterable[Sequence[object]]) -> list[dict[str, Any]]:
    """Item rows of one sheet with the family of the block they belong to."""
    items: list[dict[str, Any]] = []
    family: str | None = None
    in_table = False
    for number, cells in enumerate(rows, start=1):
        heading = _family_heading(cells)
        if heading is not None:
            # A new family keeps the sheet's columns: some blocks do not repeat the heading row.
            family = heading
        elif _is_heading_row(cells):
            in_table = True
        elif in_table and len(cells) > DESCRIPTION_COLUMN and cells[DESCRIPTION_COLUMN]:
            price = _price(cells[PRICE_COLUMN]) if len(cells) > PRICE_COLUMN else None
            items.append(
                {
                    "sheet": sheet,
                    "row": number,
                    "family": family,
                    "part_number": clean_part_number(cells[PART_COLUMN]),
                    "description": normalize_text(cells[DESCRIPTION_COLUMN]),
                    "price": price,
                    "has_price": price is not None and price > 0,  # rule R403
                    "extra_cells": _extra_cells(cells),
                }
            )
    return items


def extract_price_list(path: Path, source_sha256: str) -> pd.DataFrame:
    """Every item of every sheet; repeated part numbers flagged (rule R405)."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        items = [
            item
            for sheet in workbook.worksheets
            for item in sheet_items(sheet.title, sheet.iter_rows(values_only=True))
        ]
    finally:
        workbook.close()
    frame = pd.DataFrame(
        items,
        columns=[
            c for c in COLUMNS if c not in {"item_id", "source_sha256", "part_number_repeated"}
        ],
    )
    frame.insert(0, "source_sha256", source_sha256)
    frame.insert(0, "item_id", [f"P{i:04d}" for i in range(1, len(frame) + 1)])
    known = frame["part_number"].dropna()
    frame["part_number_repeated"] = frame["part_number"].isin(known[known.duplicated()])
    return frame[COLUMNS]
